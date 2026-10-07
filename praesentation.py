"""Stufe 3: Erstellt aus denselben Zahlen wie das Dashboard die PowerPoint ausgabe/praesentation.pptx.

Starten im Terminal:  python praesentation.py

- Liegt im Ordner vorlage/ eine eigene Firmenvorlage (.pptx), baut die Präsentation
  auf deren Master auf: Schrift, Farben, Fußzeile und Titelbild kommen von dort.
  Der Ordner vorlage/ wird nie auf GitHub hochgeladen.
- Ohne Vorlage entsteht dieselbe Präsentation in neutralem Design.

Alles bleibt bearbeitbar: echte Textfelder, echte PowerPoint-Diagramme und Tabellen,
Erklärungen in den Notizen. Die Zahlen kommen ausschließlich aus bericht.py.
"""

import sys
from datetime import datetime
from io import BytesIO
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_TICK_LABEL_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Emu, Inches, Pt

from bericht import Datenfehler, berechnen, je_projekt, laden, monat_name, veraenderung, zahl

ORDNER = Path(__file__).parent
AUSGABE = ORDNER / "ausgabe" / "praesentation.pptx"
VORLAGEN_ORDNER = ORDNER / "vorlage"

# Farben
SCHWARZ = RGBColor(0x00, 0x00, 0x00)
TEXT = RGBColor(0x40, 0x3F, 0x45)
LEISE = RGBColor(0x73, 0x72, 0x78)
TUERKIS = RGBColor(0x3A, 0xB4, 0xBA)
TUERKIS_DUNKEL = RGBColor(0x23, 0x99, 0xA1)
BLAU = RGBColor(0x32, 0x4E, 0x98)
GRAU = RGBColor(0xB5, 0xB4, 0xBA)
HELLGRAU = RGBColor(0xEC, 0xEB, 0xF0)
ROT = RGBColor(0xD5, 0x00, 0x1C)
WEISS = RGBColor(0xFF, 0xFF, 0xFF)

# Raster in Zoll (16:9). Gilt mit und ohne Vorlage.
RAND = 0.41
BREITE = 12.52
OBEN = 1.55       # Beginn der Inhaltsfläche
UNTEN = 6.40      # Ende der Inhaltsfläche
SEITE_X = 9.55    # Beginn der schmalen Erklärspalte rechts
SEITE_B = RAND + BREITE - SEITE_X


# ---------------------------------------------------------------------------
# Grundgerüst: Folien mit oder ohne Firmenvorlage
# ---------------------------------------------------------------------------

class Rahmen:
    """Erzeugt Titel- und Inhaltsfolien. Mit Vorlage aus deren Layouts, sonst schlicht."""

    def __init__(self):
        vorlagen = sorted(VORLAGEN_ORDNER.glob("*.pptx")) if VORLAGEN_ORDNER.exists() else []
        self.vorlage = vorlagen[0] if vorlagen else None
        if self.vorlage:
            self.praes = Presentation(self.vorlage)
            self.titelbild = self._erstes_bild()
            self._beispielfolien_entfernen()
            self.layout_titel = self._layout_titelfolie()
            self.layout_inhalt = self._layout_inhaltsfolie()
        else:
            self.praes = Presentation()
            self.praes.slide_width, self.praes.slide_height = Inches(13.333), Inches(7.5)
            self.titelbild = None
            self.layout_titel = self.layout_inhalt = self.praes.slide_layouts[5]  # nur Titel

    # --- Vorlage auswerten -------------------------------------------------

    def _erstes_bild(self):
        """Nimmt das Titelbild aus der ersten Folie der Vorlage, falls vorhanden."""
        if not len(self.praes.slides):
            return None
        for form in self.praes.slides[0].shapes:
            try:
                return form.image.blob
            except (AttributeError, ValueError):
                continue
        return None

    def _beispielfolien_entfernen(self):
        liste = self.praes.slides._sldIdLst
        for eintrag in list(liste):
            liste.remove(eintrag)
            self.praes.part.drop_rel(eintrag.rId)

    def _layout_titelfolie(self):
        def arten(layout):
            return {p.placeholder_format.type for p in layout.placeholders}
        kandidaten = [l for l in self.praes.slide_layouts if PP_PLACEHOLDER.CENTER_TITLE in arten(l)]
        mit_bild = [l for l in kandidaten if PP_PLACEHOLDER.PICTURE in arten(l)]
        if self.titelbild and mit_bild:
            return mit_bild[0]
        ohne_bild = [l for l in kandidaten if PP_PLACEHOLDER.PICTURE not in arten(l)]
        return (ohne_bild or kandidaten or [self.praes.slide_layouts[0]])[0]

    def _layout_inhaltsfolie(self):
        """Layout mit Titel und höchstens einer kleinen Fußnotenzeile."""
        passend = []
        for layout in self.praes.slide_layouts:
            arten = [p.placeholder_format.type for p in layout.placeholders]
            andere = [p for p in layout.placeholders if p.placeholder_format.type != PP_PLACEHOLDER.TITLE]
            if PP_PLACEHOLDER.TITLE in arten and all(Emu(p.height).inches < 0.5 for p in andere):
                passend.append((len(andere), layout))
        if passend:
            return sorted(passend, key=lambda e: e[0])[0][1]
        return self.praes.slide_layouts[0]

    # --- Folien anlegen ----------------------------------------------------

    def titelfolie(self, titel, untertitel):
        folie = self.praes.slides.add_slide(self.layout_titel)
        if self.vorlage:
            for ph in list(folie.placeholders):
                art = ph.placeholder_format.type
                if art in (PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.SUBTITLE):
                    ph.text = titel if art != PP_PLACEHOLDER.SUBTITLE else untertitel
                    if self.titelbild:  # auf dem Foto: ohne Fläche, weiße Schrift
                        ph.fill.background()
                        for absatz in ph.text_frame.paragraphs:
                            for lauf in absatz.runs:
                                lauf.font.color.rgb = WEISS
                elif art == PP_PLACEHOLDER.PICTURE and self.titelbild:
                    ph.insert_picture(BytesIO(self.titelbild))
                else:
                    ph._element.getparent().remove(ph._element)
        else:
            folie.shapes.title._element.getparent().remove(folie.shapes.title._element)
            flaeche = folie.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(4.2), self.praes.slide_height)
            verlauf(flaeche)
            textfeld(folie, 4.9, 2.55, 7.8, 1.2, titel, groesse=48, farbe=SCHWARZ, fett=True)
            textfeld(folie, 4.9, 3.75, 7.8, 0.6, untertitel, groesse=18, farbe=TEXT)
        return folie

    def inhaltsfolie(self, titel, unterzeile, quelle):
        """Neue Inhaltsfolie: Titel als Aussage, darunter eine kurze Einordnung, unten die Quelle."""
        folie = self.praes.slides.add_slide(self.layout_inhalt)
        titelfeld = folie.shapes.title
        titelfeld.text = titel
        if not self.vorlage:
            titelfeld.left, titelfeld.top = Inches(RAND), Inches(0.42)
            titelfeld.width, titelfeld.height = Inches(BREITE), Inches(0.71)
            absatz = titelfeld.text_frame.paragraphs[0]
            absatz.alignment = PP_ALIGN.LEFT
            absatz.font.size, absatz.font.bold, absatz.font.color.rgb = Pt(28), True, SCHWARZ
            titelfeld.text_frame.vertical_anchor = MSO_ANCHOR.BOTTOM
        beschriftung(folie, RAND, 1.13, BREITE, 0.3, unterzeile)

        fussnote = [p for p in folie.placeholders if p.placeholder_format.type != PP_PLACEHOLDER.TITLE]
        if fussnote:
            fussnote[0].text = quelle
            for p in fussnote[1:]:
                p._element.getparent().remove(p._element)
        else:
            textfeld(folie, RAND, 6.52, BREITE, 0.25, quelle, groesse=9, farbe=LEISE)
        return folie


# ---------------------------------------------------------------------------
# Bausteine für Folieninhalte (auch für eigene Folien verwenden)
# ---------------------------------------------------------------------------

def textfeld(folie, links, oben, breite, hoehe, text, groesse=14, farbe=TEXT, fett=False,
             ausrichtung=PP_ALIGN.LEFT, anker=MSO_ANCHOR.TOP, sperrung=0):
    """Ein bearbeitbares Textfeld. Zeilenumbrüche im Text werden zu eigenen Absätzen."""
    feld = folie.shapes.add_textbox(Inches(links), Inches(oben), Inches(breite), Inches(hoehe))
    rahmen = feld.text_frame
    rahmen.word_wrap = True
    rahmen.vertical_anchor = anker
    rahmen.margin_left = rahmen.margin_right = rahmen.margin_top = rahmen.margin_bottom = 0
    for i, zeile in enumerate(str(text).split("\n")):
        absatz = rahmen.paragraphs[0] if i == 0 else rahmen.add_paragraph()
        absatz.alignment = ausrichtung
        lauf = absatz.add_run()
        lauf.text = zeile
        lauf.font.size, lauf.font.bold, lauf.font.color.rgb = Pt(groesse), fett, farbe
        if sperrung:
            lauf.font._element.set("spc", str(sperrung))
    return feld


def beschriftung(folie, links, oben, breite, hoehe, text, farbe=TUERKIS):
    """Kleine Überschrift in gesperrten Großbuchstaben, z. B. AUF EINEN BLICK."""
    return textfeld(folie, links, oben, breite, hoehe, text.upper(), groesse=10.5, farbe=farbe,
                    fett=True, sperrung=150)


def linie(folie, links, oben, breite, farbe=HELLGRAU, staerke=0.75):
    verbinder = folie.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(links), Inches(oben),
                                           Inches(links + breite), Inches(oben))
    verbinder.line.color.rgb = farbe
    verbinder.line.width = Pt(staerke)
    return verbinder


def verlauf(form):
    """Fläche mit Verlauf von Blau nach Türkis, ohne Rand."""
    form.fill.gradient()
    form.fill.gradient_angle = 45
    stopps = form.fill.gradient_stops
    stopps[0].color.rgb, stopps[0].position = BLAU, 0.0
    stopps[1].color.rgb, stopps[1].position = TUERKIS_DUNKEL, 1.0
    form.line.fill.background()
    form.shadow.inherit = False
    return form


def verlaufsflaeche(folie, links, oben, breite, hoehe):
    return verlauf(folie.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(links), Inches(oben),
                                          Inches(breite), Inches(hoehe)))


def erklaerspalte(folie, titel, absaetze, oben=OBEN):
    """Schmale Spalte rechts: wie man die Folie liest."""
    linie(folie, SEITE_X, oben, SEITE_B, farbe=TUERKIS, staerke=1.5)
    beschriftung(folie, SEITE_X, oben + 0.15, SEITE_B, 0.3, titel)
    feld = textfeld(folie, SEITE_X, oben + 0.55, SEITE_B, UNTEN - oben - 0.6, "\n".join(absaetze),
                    groesse=12, farbe=TEXT)
    for absatz in feld.text_frame.paragraphs:
        absatz.space_after = Pt(9)
    return feld


def tabelle(folie, kopf, zeilen, links, oben, breite, spaltenbreiten=None, rechtsbuendig=(),
            zeilenhoehe=0.36, schrift=12, streifen=True):
    """Bearbeitbare PowerPoint-Tabelle im Stil der Präsentation."""
    form = folie.shapes.add_table(len(zeilen) + 1, len(kopf), Inches(links), Inches(oben),
                                  Inches(breite), Inches(zeilenhoehe * (len(zeilen) + 1)))
    tab = form.table
    tab.first_row, tab.horz_banding = True, False
    anteile = spaltenbreiten or [1] * len(kopf)
    for i, anteil in enumerate(anteile):
        tab.columns[i].width = Inches(breite * anteil / sum(anteile))
    for r in range(len(zeilen) + 1):
        tab.rows[r].height = Inches(zeilenhoehe)
        for c in range(len(kopf)):
            zelle = tab.cell(r, c)
            zelle.text = str(kopf[c] if r == 0 else zeilen[r - 1][c])
            zelle.vertical_anchor = MSO_ANCHOR.MIDDLE
            zelle.margin_left = zelle.margin_right = Inches(0.1)
            zelle.margin_top = zelle.margin_bottom = Inches(0.03)
            zelle.fill.solid()
            zelle.fill.fore_color.rgb = BLAU if r == 0 else (HELLGRAU if r % 2 == 0 or not streifen else WEISS)
            absatz = zelle.text_frame.paragraphs[0]
            absatz.alignment = PP_ALIGN.RIGHT if c in rechtsbuendig else PP_ALIGN.LEFT
            absatz.font.size = Pt(schrift)
            absatz.font.bold = r == 0
            absatz.font.color.rgb = WEISS if r == 0 else TEXT
    return form


def balkendiagramm(folie, kategorien, reihen, links, oben, breite, hoehe, beschriftete_reihe=None,
                   zahlenformat='0" h"', legende=True, senkrecht=False):
    """Balkendiagramm (echtes PowerPoint-Diagramm). reihen = [(Name, Werte, Farbe), ...].

    Waagerecht für Listen wie Projekte, mit senkrecht=True als Säulen, z. B. für Monate.
    """
    daten = CategoryChartData()
    daten.categories = kategorien
    for name, werte, _ in reihen:
        daten.add_series(name, [float(w) for w in werte])
    art = XL_CHART_TYPE.COLUMN_CLUSTERED if senkrecht else XL_CHART_TYPE.BAR_CLUSTERED
    diagramm = folie.shapes.add_chart(art, Inches(links), Inches(oben),
                                      Inches(breite), Inches(hoehe), daten).chart
    diagramm.has_title = False
    diagramm.font.size = Pt(11)
    diagramm.font.color.rgb = TEXT
    plot = diagramm.plots[0]
    plot.gap_width = 55
    plot.overlap = -5 if len(reihen) > 1 else 0
    for serie, (_, _, farbe) in zip(plot.series, reihen):
        serie.invert_if_negative = False
        serie.format.fill.solid()
        serie.format.fill.fore_color.rgb = farbe
        serie.format.line.fill.background()
    if beschriftete_reihe is not None:
        etiketten = plot.series[beschriftete_reihe].data_labels
        etiketten.number_format, etiketten.number_format_is_linked = zahlenformat, False
        etiketten.position = XL_LABEL_POSITION.OUTSIDE_END
        etiketten.font.size, etiketten.font.color.rgb = Pt(10), TEXT
        etiketten.show_value = True
    kategorie_achse = diagramm.category_axis
    kategorie_achse.reverse_order = not senkrecht  # Balken: erste Kategorie oben
    kategorie_achse.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    kategorie_achse.format.line.color.rgb = GRAU
    kategorie_achse.has_major_gridlines = False
    wert_achse = diagramm.value_axis
    if all(float(w) >= 0 for _, werte, _ in reihen for w in werte):
        wert_achse.minimum_scale = 0  # Balken immer ab 0, sonst wirken Unterschiede größer als sie sind
    wert_achse.visible = False
    wert_achse.has_major_gridlines = False
    diagramm.has_legend = legende and len(reihen) > 1
    if diagramm.has_legend:
        diagramm.legend.position = XL_LEGEND_POSITION.TOP
        diagramm.legend.include_in_layout = False
        diagramm.legend.font.size = Pt(11)
    return diagramm


def punkte_einfaerben(diagramm, reihe, farben):
    """Färbt einzelne Balken einer Reihe, z. B. rot, wenn Ist über Plan liegt."""
    serie = diagramm.plots[0].series[reihe]
    for i, farbe in enumerate(farben):
        punkt = serie.points[i]
        punkt.format.fill.solid()
        punkt.format.fill.fore_color.rgb = farbe
        # Negative Werte sonst weiß dargestellt: Invertieren je Balken ausschalten
        dpt = serie._element.get_or_add_dPt_for_point(i)
        if dpt.find(qn("c:invertIfNegative")) is None:
            element = OxmlElement("c:invertIfNegative")
            element.set("val", "0")
            dpt.find(qn("c:idx")).addnext(element)


def notizen(folie, text):
    folie.notes_slide.notes_text_frame.text = text


# ---------------------------------------------------------------------------
# Die Folien
# ---------------------------------------------------------------------------

def folie_titel(rahmen, df, k):
    folie = rahmen.titelfolie(
        "Projektstatus",
        f"{k['zeitraum']} · {k['anzahl_projekte']} Projekte",
    )
    notizen(folie, "Automatisch erstellt aus allen Monatsdateien im Ordner daten/. Alle Zahlen stammen aus derselben "
                   "Berechnung wie das Dashboard. Die Folie „Einordnung und nächste Schritte“ ergänzt ihr selbst.")


def folie_ueberblick(rahmen, df, k, quelle):
    anteil = k["ausschoepfung"]
    titel = (f"{zahl(anteil * 100)} % des geplanten Aufwands sind verbraucht" if anteil is not None
             else "Kennzahlen auf einen Blick")
    folie = rahmen.inhaltsfolie(titel, "Auf einen Blick", quelle)

    # Links: die wichtigste Zahl groß auf der Verlaufsfläche
    verlaufsflaeche(folie, RAND, OBEN, 4.55, UNTEN - OBEN)
    beschriftung(folie, RAND + 0.45, OBEN + 0.45, 3.7, 0.3, "Ausschöpfung", farbe=WEISS)
    textfeld(folie, RAND + 0.45, OBEN + 0.85, 3.7, 1.3,
             f"{zahl(anteil * 100)} %" if anteil is not None else "–", groesse=66, farbe=WEISS, fett=True)
    textfeld(folie, RAND + 0.45, OBEN + 2.35, 3.7, 0.9,
             f"{zahl(k['ist_stunden'])} von {zahl(k['plan_stunden'])} geplanten Stunden sind gebucht.",
             groesse=16, farbe=WEISS)
    # Fortschrittsbalken: dünne weiße Linie = 100 % des Plans, kräftiger Balken = verbraucht
    balken_y = UNTEN - 0.75
    linie(folie, RAND + 0.45, balken_y + 0.05, 3.65, farbe=WEISS, staerke=0.75)
    if anteil is not None:
        fuellung = folie.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(RAND + 0.45), Inches(balken_y),
                                          Inches(3.65 * min(anteil, 1)), Inches(0.1))
        fuellung.fill.solid()
        fuellung.fill.fore_color.rgb = ROT if anteil > 1 else WEISS
        fuellung.line.fill.background()
        fuellung.shadow.inherit = False
        textfeld(folie, RAND + 0.45, balken_y + 0.2, 3.65, 0.3, "0 %", groesse=10, farbe=WEISS)
        textfeld(folie, RAND + 0.45, balken_y + 0.2, 3.65, 0.3, "100 % des Plans", groesse=10, farbe=WEISS,
                 ausrichtung=PP_ALIGN.RIGHT)

    # Rechts: weitere Kennzahlen als Liste mit Trennlinien
    zeilen = [
        (zahl(k["anzahl_projekte"]), "Projekte", k["zeitraum"]),
        (f"{zahl(k['plan_stunden'])} h", "Plan-Stunden", "Summe der geplanten Stunden"),
        (f"{zahl(k['ist_stunden'])} h", "Ist-Stunden", "Summe der gebuchten Stunden"),
        (f"{zahl(k['abweichung'], vorzeichen=True)} h", "Abweichung",
         "Ist minus Plan · negativ heißt: weniger verbraucht als geplant"),
    ]
    x, hoehe = RAND + 5.05, (UNTEN - OBEN) / len(zeilen)
    for i, (wert, name, erklaerung) in enumerate(zeilen):
        y = OBEN + i * hoehe
        if i:
            linie(folie, x, y, RAND + BREITE - x)
        textfeld(folie, x, y + 0.18, 3.0, 0.8, wert, groesse=32, farbe=SCHWARZ, fett=True, anker=MSO_ANCHOR.MIDDLE)
        beschriftung(folie, x + 3.2, y + 0.26, 4.2, 0.3, name)
        textfeld(folie, x + 3.2, y + 0.58, 4.2, 0.5, erklaerung, groesse=12, farbe=LEISE)
    notizen(folie, "Ausschöpfung = Ist-Stunden geteilt durch Plan-Stunden. Unter 100 % wurde weniger "
                   "Aufwand gebucht als geplant. Das ist nicht automatisch gut: Projekte, die noch nicht "
                   "gestartet sind, senken den Wert ebenfalls (siehe Folie „Plan und Ist je Projekt“).")


def folie_entwicklung(rahmen, df, k, quelle):
    """Plan und Ist je Monat: zeigt, wie sich der Aufwand mit jeder Monatsdatei entwickelt."""
    monate = sorted(df["Monat"].unique())
    werte = [berechnen(df[df["Monat"] == m]) for m in monate]
    titel = f"Im {monat_name(monate[-1])} wurden {zahl(werte[-1]['ist_stunden'])} Stunden gebucht"
    if len(monate) > 1:
        differenz = veraenderung(werte[-1], werte[-2])["ist_stunden"]
        if differenz:
            titel += (f", {zahl(abs(differenz))} {'mehr' if differenz > 0 else 'weniger'} "
                      f"als im {monat_name(monate[-2]).split(' ')[0]}")
    folie = rahmen.inhaltsfolie(titel, "Entwicklung je Monat", quelle)

    namen = [monat_name(m) for m in monate]
    diagramm = balkendiagramm(
        folie, namen, [("Plan", [w["plan_stunden"] for w in werte], GRAU),
                       ("Ist", [w["ist_stunden"] for w in werte], BLAU)],
        RAND, OBEN - 0.05, 6.6, UNTEN - OBEN + 0.15, beschriftete_reihe=1, senkrecht=True,
    )
    punkte_einfaerben(diagramm, 1, [ROT if w["ist_stunden"] > w["plan_stunden"] else BLAU for w in werte])

    zeilen = []
    for i, (name, w) in enumerate(zip(namen, werte)):
        vorher = zahl(veraenderung(w, werte[i - 1])["ist_stunden"], vorzeichen=True) if i else "–"
        zeilen.append((name, zahl(w["plan_stunden"]), zahl(w["ist_stunden"]), vorher))
    tabelle(folie, ["Monat", "Plan h", "Ist h", "Ist ggü. Vormonat h"], zeilen, RAND + 7.0, OBEN + 0.35,
            BREITE - 7.0, spaltenbreiten=[2.2, 1, 1, 1.8], rechtsbuendig={1, 2, 3})
    unter_tabelle = OBEN + 0.35 + 0.36 * (len(zeilen) + 1) + 0.5
    linie(folie, RAND + 7.0, unter_tabelle, BREITE - 7.0, farbe=TUERKIS, staerke=1.5)
    beschriftung(folie, RAND + 7.0, unter_tabelle + 0.15, BREITE - 7.0, 0.3, "So lesen")
    textfeld(folie, RAND + 7.0, unter_tabelle + 0.55, BREITE - 7.0, 1.2,
             "Jede Säule ist eine Monatsdatei aus dem Ordner daten/. Kommt eine neue Datei dazu, "
             "erscheint automatisch ein weiterer Monat.\nRot: In diesem Monat wurde mehr gebucht als geplant.",
             groesse=12)
    notizen(folie, "Plan und Ist sind die Stunden des jeweiligen Monats, nicht aufsummiert. "
                   "Die Veränderung vergleicht die Ist-Stunden mit dem Monat davor.")


def folie_plan_ist(rahmen, df, k, quelle):
    sortiert = je_projekt(df).sort_values("Plan_Stunden", ascending=False)
    folie = rahmen.inhaltsfolie("Plan und Ist je Projekt im Vergleich", f"Plan und Ist · {k['zeitraum']}", quelle)
    diagramm = balkendiagramm(
        folie, list(sortiert["Projekt"]),
        [("Plan", sortiert["Plan_Stunden"], GRAU), ("Ist", sortiert["Ist_Stunden"], BLAU)],
        RAND, OBEN - 0.05, SEITE_X - RAND - 0.35, UNTEN - OBEN + 0.15, beschriftete_reihe=1,
    )
    punkte_einfaerben(diagramm, 1, [ROT if i > p else BLAU
                                    for i, p in zip(sortiert["Ist_Stunden"], sortiert["Plan_Stunden"])])
    erklaerspalte(folie, "So lesen", [
        "Grau: geplante Stunden.",
        "Blau: gebuchte Stunden, mit Wert am Balkenende.",
        "Rot: gebucht ist mehr als geplant.",
        "Stunden je Projekt über alle Monate addiert, größte Projekte oben.",
    ])
    notizen(folie, "Echtes PowerPoint-Diagramm: Rechtsklick → Daten bearbeiten zeigt die Werte. "
                   "Rote Balken markieren Projekte, bei denen Ist größer als Plan ist. Ein Projekt ohne "
                   "gebuchte Stunden ist meist noch nicht gestartet: Status in der Tabelle prüfen.")


def folie_bereiche(rahmen, df, k, quelle):
    summen = (df.groupby("Bereich")[["Plan_Stunden", "Ist_Stunden"]].sum()
              .sort_values("Ist_Stunden", ascending=False))
    groesster = summen.index[0]
    folie = rahmen.inhaltsfolie(
        f"{groesster} bindet mit {zahl(summen.iloc[0]['Ist_Stunden'])} Stunden den größten Aufwand",
        "Stunden je Bereich", quelle)
    balkendiagramm(
        folie, list(summen.index),
        [("Plan", summen["Plan_Stunden"], GRAU), ("Ist", summen["Ist_Stunden"], BLAU)],
        RAND, OBEN - 0.05, 6.6, UNTEN - OBEN + 0.15, beschriftete_reihe=1,
    )
    zeilen = [(b, zahl(z["Plan_Stunden"]), zahl(z["Ist_Stunden"]),
               zahl(z["Ist_Stunden"] - z["Plan_Stunden"], vorzeichen=True)) for b, z in summen.iterrows()]
    tabelle(folie, ["Bereich", "Plan h", "Ist h", "Abw. h"], zeilen, RAND + 7.0, OBEN + 0.35,
            BREITE - 7.0, spaltenbreiten=[2.2, 1, 1, 1], rechtsbuendig={1, 2, 3})
    notizen(folie, "Summen der Plan- und Ist-Stunden aller Projekte je Bereich. Sortiert nach Ist-Stunden.")


def folie_projekte(rahmen, df, k, quelle):
    """Alle Projekte als Tabelle. Bei vielen Projekten auf mehrere Folien verteilt."""
    zeilen = [(z["Projekt"], z["Bereich"], z["Status"], zahl(z["Plan_Stunden"]), zahl(z["Ist_Stunden"]),
               zahl(z["Ist_Stunden"] - z["Plan_Stunden"], vorzeichen=True)) for _, z in je_projekt(df).iterrows()]
    pro_folie = 11
    teile = [zeilen[i:i + pro_folie] for i in range(0, len(zeilen), pro_folie)]
    for nr, teil in enumerate(teile, start=1):
        zusatz = f" ({nr}/{len(teile)})" if len(teile) > 1 else ""
        folie = rahmen.inhaltsfolie(f"Alle Projekte im Detail{zusatz}", "Datengrundlage", quelle)
        tabelle(folie, ["Projekt", "Bereich", "Status (zuletzt)", "Plan h", "Ist h", "Abweichung h"], teil,
                RAND, OBEN, BREITE, spaltenbreiten=[3, 2, 2, 1.2, 1.2, 1.5], rechtsbuendig={3, 4, 5},
                zeilenhoehe=0.38, schrift=13)
        notizen(folie, "Alle Projekte, Stunden über alle Monate addiert, Status aus dem letzten Monat. Werte sind Text in einer echten "
                       "Tabelle und können für Ergänzungen bearbeitet werden. Die Quelldatei bleibt maßgeblich.")


def folie_ueber_plan(rahmen, df, k, quelle):
    projekte = je_projekt(df)  # jedes Projekt über alle Monate zusammengefasst
    ueber = projekte[projekte["Ist_Stunden"] > projekte["Plan_Stunden"]]
    ueber = ueber.assign(Abweichung=ueber["Ist_Stunden"] - ueber["Plan_Stunden"]).sort_values("Abweichung", ascending=False)
    anzahl = k["projekte_ueber_plan"]
    folie = rahmen.inhaltsfolie(f"{anzahl} {'Projekt liegt' if anzahl == 1 else 'Projekte liegen'} über Plan",
                                "Ist größer als Plan", quelle)
    if ueber.empty:
        textfeld(folie, RAND, OBEN, BREITE, 0.6, "Kein Projekt liegt über Plan.", groesse=18)
    else:
        zeilen = [(z["Projekt"], z["Bereich"], zahl(z["Plan_Stunden"]), zahl(z["Ist_Stunden"]),
                   zahl(z["Abweichung"], vorzeichen=True)) for _, z in ueber.iterrows()]
        tabelle(folie, ["Projekt", "Bereich", "Plan h", "Ist h", "Abweichung h"], zeilen,
                RAND, OBEN, SEITE_X - RAND - 0.45, spaltenbreiten=[3, 2, 1.2, 1.2, 1.5], rechtsbuendig={2, 3, 4},
                zeilenhoehe=0.42, schrift=14)
    erklaerspalte(folie, "So lesen", [
        f"Projekte, bei denen im Zeitraum {k['zeitraum']} mehr Stunden gebucht als geplant sind.",
        "Sortiert nach der größten Überschreitung.",
        "Ursachen gehören auf die Folie „Einordnung und nächste Schritte“.",
    ])
    notizen(folie, "Kennzahl „Projekte über Plan“ aus bericht.py: Anzahl Projekte, deren Ist-Stunden über alle "
                   "Monate zusammen größer sind als die Plan-Stunden.")


def folie_aussage(rahmen, df, k, quelle):
    folie = rahmen.inhaltsfolie("Einordnung und nächste Schritte", "Von euch zu ergänzen", quelle)
    breite_links = 5.6
    beschriftung(folie, RAND, OBEN, breite_links, 0.3, "Einordnung")
    fragen = [
        ("Was fällt auf?", "[Die zwei bis drei wichtigsten Beobachtungen]"),
        ("Woran liegt es?", "[Nur mit Beleg, sonst als offene Frage formulieren]"),
        ("Was empfehlen wir?", "[Konkrete Empfehlung]"),
    ]
    y = OBEN + 0.45
    for frage, platzhalter in fragen:
        linie(folie, RAND, y, breite_links)
        textfeld(folie, RAND, y + 0.15, breite_links, 0.35, frage, groesse=16, farbe=SCHWARZ, fett=True)
        textfeld(folie, RAND, y + 0.55, breite_links, 0.6, platzhalter, groesse=14, farbe=LEISE)
        y += 1.45
    x = RAND + breite_links + 0.6
    beschriftung(folie, x, OBEN, RAND + BREITE - x, 0.3, "Nächste Schritte")
    tabelle(folie, ["Maßnahme", "Verantwortlich", "Bis wann"], [("", "", "")] * 5,
            x, OBEN + 0.45, RAND + BREITE - x, spaltenbreiten=[3, 1.6, 1.1], zeilenhoehe=0.62, schrift=13, streifen=False)
    notizen(folie, "Diese Folie wird bewusst nicht automatisch ausgefüllt. Die Zahlen auf den anderen Folien "
                   "zeigen, was passiert ist. Bewertung, Ursachen und Empfehlungen formuliert ein Mensch.")


def folie_methodik(rahmen, df, k, quelle):
    folie = rahmen.inhaltsfolie("Über diesen Bericht", "Datengrundlage und Definitionen", quelle)
    spalten = [
        ("Datenquelle", [
            f"Ordner: daten/ ({df['Monat'].nunique()} Monatsdateien)",
            f"Zeitraum: {k['zeitraum']}",
            f"Projekte: {k['anzahl_projekte']}",
            f"Erstellt am: {datetime.now().strftime('%d.%m.%Y %H:%M')}",
            "Erzeugt mit: praesentation.py",
        ]),
        ("Automatisch geprüft", [
            "✓ Alle Pflichtspalten vorhanden",
            "✓ Kein Projekt doppelt",
            "✓ Plan- und Ist-Stunden vollständig und nicht negativ",
            "✓ Nur erlaubte Statuswerte",
            "✓ Jede Datei enthält genau einen Monat",
            "✓ Kein Monat doppelt",
        ]),
        ("Definitionen", [
            "Abweichung = Ist-Stunden − Plan-Stunden",
            "Ausschöpfung = Ist-Stunden ÷ Plan-Stunden",
            "Je Projekt: Stunden über alle Monate addiert",
            "Status: Stand im letzten Monat",
        ]),
    ]
    breite = (BREITE - 2 * 0.5) / 3
    for i, (titel, eintraege) in enumerate(spalten):
        x = RAND + i * (breite + 0.5)
        linie(folie, x, OBEN, breite, farbe=TUERKIS, staerke=1.5)
        beschriftung(folie, x, OBEN + 0.15, breite, 0.3, titel)
        textfeld(folie, x, OBEN + 0.6, breite, 3.8, "\n".join(eintraege), groesse=13, farbe=TEXT)
    notizen(folie, "Hält fest, woher die Zahlen stammen und wie sie berechnet sind. Bei Rückfragen "
                   "zuerst den Zeitraum vergleichen.")


def main():
    try:
        df = laden()
    except Datenfehler as fehler:
        print("FEHLER: Mindestens eine Excel-Datei ist nicht in Ordnung. Es wurde keine Präsentation erstellt.\n")
        print(fehler)
        sys.exit(1)

    k = berechnen(df)
    rahmen = Rahmen()
    quelle = f"Quelle: {df['Monat'].nunique()} Monatsdateien in daten/ · {k['zeitraum']} · automatisch erstellt"

    folie_titel(rahmen, df, k)
    folie_ueberblick(rahmen, df, k, quelle)
    folie_entwicklung(rahmen, df, k, quelle)
    folie_plan_ist(rahmen, df, k, quelle)
    folie_bereiche(rahmen, df, k, quelle)
    folie_projekte(rahmen, df, k, quelle)
    folie_ueber_plan(rahmen, df, k, quelle)
    folie_aussage(rahmen, df, k, quelle)
    folie_methodik(rahmen, df, k, quelle)

    AUSGABE.parent.mkdir(exist_ok=True)
    rahmen.praes.save(AUSGABE)
    art = f"Vorlage {rahmen.vorlage.name}" if rahmen.vorlage else "neutrales Design"
    print(f"Praesentation erstellt: {AUSGABE.name} ({len(rahmen.praes.slides)} Folien, {art}, "
          f"{k['zeitraum']})")


if __name__ == "__main__":
    main()
